"""Import d'un rapport d'inspection Excel « par zones » (format papier des
clients industriels, ex. Rio Tinto QIT) : une feuille par zone, des lignes
de titre de section, puis une ligne par extincteur.

Colonnes d'une feuille de zone (en-tête « Numéro / Number » en colonne A) :
  A n° de ligne   B numéro d'inventaire   C emplacement   D code équipement
  E modèle        F année fabrication     G n° série      H marque
  I dernier hydro J prochain hydro        K prochain entretien
  L non-conformités   M à faire   N fait
"""
import re
import unicodedata

from openpyxl import load_workbook

from .models import CODES_NON_CONFORMITES, ExtincteurItem

MAX_LIGNES_ENTETE = 40


def _texte(v) -> str:
    if v is None or isinstance(v, bool):
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip()


def _sans_accents(t: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", t) if unicodedata.category(c) != "Mn").lower()


def _annee(v):
    """Année à 4 chiffres, ou None (« inc », « - », « . », vide…)."""
    if isinstance(v, (int, float)) and 1900 <= int(v) <= 2100:
        return int(v)
    m = re.search(r"\b(19|20)\d{2}\b", _texte(v))
    return int(m.group(0)) if m else None


def _valeur_inconnue(t: str) -> bool:
    return _sans_accents(t).strip(" .") in ("", "-", "inc", "onc", "?", "x", "na", "n/a")


# ── Modèle → type + format ──────────────────────────────────────────────────
def _type_format(modele: str) -> tuple[str, str]:
    t = _sans_accents(modele).replace("c02", "co2").replace(",", ".")
    if re.fullmatch(r"(19|20)\d{2}", t.strip()):
        return "", ""  # année saisie par erreur dans la colonne modèle
    T = ExtincteurItem.TypeExtincteur
    if "abc" in t:
        type_ = T.POUDRE_ABC
    elif re.search(r"\bbc\b|^bc|\bbc\d", t):
        type_ = T.POUDRE_BC
    elif "co2" in t:
        type_ = T.CO2
    elif "eau" in t or "water" in t:
        type_ = T.EAU
    elif "type k" in t or re.search(r"\bk\b", t):
        type_ = T.K
    elif "fe 36" in t or "fe36" in t:
        type_ = T.FE36
    elif re.search(r"\bpk\b", t):
        type_ = T.PK
    elif "classe d" in t or re.search(r"(^|\s)d\s*\d", t):
        type_ = T.D
    else:
        type_ = ""

    format_ = ""
    if type_ != T.FE36:
        m = re.search(r"(\d+(?:\.\d+)?)", t.replace("co2", ""))
        if m:
            format_ = {
                "2.5": "2.5lb", "5": "5lb", "10": "10lb", "15": "15lb", "20": "20lb",
                "30": "30lb", "50": "50lb", "125": "125lb",
            }.get(m.group(1).rstrip("0").rstrip(".") if "." in m.group(1) else m.group(1), "")
    return type_, format_


# ── Marque ──────────────────────────────────────────────────────────────────
_MARQUES = [
    (r"str?i?ke?[\s-]*fi+rst|^sf$|^st-?first$", "strikefirst"),
    (r"ansul", "ansul"),
    (r"general", "general"),
    (r"flag", "flag"),
    (r"amerex", "amerex"),
    (r"buck", "buckeye"),
    (r"kidde", "kidde"),
    (r"badger|bager", "badger"),
    (r"pyr+en+e|purene", "pyrene"),
    (r"sentry", "sentry"),
    (r"pyro\s*c", "pyrochem"),
]


def _marque(v) -> str:
    t = _sans_accents(_texte(v))
    if _valeur_inconnue(t) or re.fullmatch(r"\d{4}", t):
        return ""
    for motif, code in _MARQUES:
        if re.search(motif, t):
            return code
    return "autre"


# ── Non-conformités (texte libre → codes de la légende + remarque) ──────────
_ALIAS_NC = {"MISS": "MIS", "MANQUANT": "MIS", "6ANS": "6Y", "SCRAP": "RP", "REMPLACER": "RP"}


def _non_conformites(v) -> tuple[list[str], str]:
    texte = _texte(v)
    if not texte or re.fullmatch(r"(?i)ok[\s,.]*(new|neuf)?[\s\d/-]*", texte):
        return [], ""  # « OK », « Ok, new », « OK NEW 04-2021 » : rien à signaler
    codes = []
    for mot in re.findall(r"[A-Za-z0-9]+", texte):
        code = _ALIAS_NC.get(mot.upper(), mot.upper())
        if code in CODES_NON_CONFORMITES and code not in codes:
            codes.append(code)
    reste = texte
    if codes and re.fullmatch(r"(?i)[\s,/;-]*((" + "|".join(
        [re.escape(c) for c in CODES_NON_CONFORMITES] + list(_ALIAS_NC)) + r")[\s,/;-]*)+", texte):
        reste = ""  # le texte n'était que des codes
    return [c for c in CODES_NON_CONFORMITES if c in codes], reste


# ── Sections et notes d'accès ───────────────────────────────────────────────
_MOTS_NOTE = re.compile(
    r"\b(permis|aviser|avertir|demande|pas besoin|note|inscription|il y a|toujours|acces|appel|"
    r"radio|le compte|accompagn|registre|signe)", re.I
)


def _est_note(t: str) -> bool:
    return t.startswith("(") or bool(_MOTS_NOTE.search(_sans_accents(t)))


def _separer_nom_note(t: str) -> tuple[str, str]:
    """« ATELIER MÉCANIQUE ( toujours aviser M. Coté) » → nom + note."""
    m = re.match(r"^(.{3,}?)\s*\((.*)$", t)
    if m and not t.startswith("("):
        note = m.group(2).strip()
        if note.endswith(")"):
            note = note[:-1].strip()  # la parenthèse ouverte par le titre
        return m.group(1).strip(" ,-"), note
    m = re.match(r"^(.{3,}?),\s*(.*)$", t)
    if m and _est_note(m.group(2)):
        return m.group(1).strip(), m.group(2).strip()
    return t, ""


def _nom_lieu(titre_feuille: str) -> str:
    nom = re.sub(r"\s*\((ext|list)\)\s*$", "", titre_feuille.strip(), flags=re.I).strip()
    nom = re.sub(r"^ZONE\b", "Zone", nom)
    return nom[:1].upper() + nom[1:]


def _ligne_entete(ws) -> int | None:
    for r in range(1, MAX_LIGNES_ENTETE + 1):
        if (_texte(ws.cell(r, 1).value).lower().startswith("numéro")
                and "emplacement" in _texte(ws.cell(r, 3).value).lower()
                and "quipement" in _texte(ws.cell(r, 4).value).lower()):
            return r
    return None


def _lire_feuille(ws, ligne_entete: int) -> dict:
    sections: list[dict] = []
    notes_en_attente: list[str] = []
    courante = None
    precedente_titre = False  # la ligne d'avant était une ligne de titre

    def nouvelle_section(nom: str, note: str):
        nonlocal courante
        # Section vide suivie d'une autre : on les fusionne
        # (« ADMINISTRATION PRINCIPALE » + « Édifice administratif »).
        if courante and not courante["extincteurs"] and precedente_titre:
            # Au plus deux niveaux : « ZONE 200 — ATELIER POOL ».
            courante["nom"] = f"{courante['nom'].split(' — ')[-1]} — {nom}"[:200]
            if note:
                courante["notes"].append(note)
            return
        courante = {"nom": nom[:200], "notes": notes_en_attente.copy() + ([note] if note else []), "extincteurs": []}
        notes_en_attente.clear()
        sections.append(courante)

    for r in range(ligne_entete + 1, ws.max_row + 1):
        v = [ws.cell(r, c).value for c in range(1, 15)]
        b, c, d, e = (_texte(x) for x in v[1:5])
        donnees = [_texte(x) for x in v[3:11]]  # D..K

        if any(donnees) or (b and c):
            if courante is None:
                nouvelle_section("Général", "")
            nc_codes, remarque = _non_conformites(v[11])
            for libelle, cellule in (("À faire", v[12]), ("Fait", v[13])):
                t = _texte(cellule)
                if t:
                    remarque = f"{remarque} · {libelle} : {t}" if remarque else f"{libelle} : {t}"
            modele = e if not _valeur_inconnue(e) else ""
            type_, format_ = _type_format(modele)
            courante["extincteurs"].append({
                "numero": "" if _valeur_inconnue(b) else b[:30],
                "emplacement": c[:200],
                "code_equipement": "" if _valeur_inconnue(d) else d[:50],
                "modele": modele[:60],
                "type_extincteur": type_,
                "format": format_,
                "date_fabrication": _annee(v[5]),
                "numero_serie": "" if _valeur_inconnue(_texte(v[6])) else _texte(v[6])[:100],
                "marque": _marque(v[7]),
                "dernier_test_hydrostatique": _annee(v[8]),
                "prochain_test_hydrostatique": _annee(v[9]),
                "prochaine_maintenance": _annee(v[10]),
                "non_conformites": nc_codes,
                "remarque": remarque[:300],
            })
            precedente_titre = False
            continue

        titre = c or b  # quelques titres sont saisis en colonne B
        if not titre or titre.startswith("*"):
            continue
        if precedente_titre and _est_note(titre) and courante:
            courante["notes"].append(titre)
        elif courante is None and _est_note(_separer_nom_note(titre)[0]):
            notes_en_attente.append(titre)
        else:
            nom, note = _separer_nom_note(titre)
            nouvelle_section(nom, note)
        precedente_titre = True

    sections = [s for s in sections if s["extincteurs"]]
    return {
        "sections": [
            {"nom": s["nom"], "note": "\n".join(s["notes"]), "extincteurs": s["extincteurs"]}
            for s in sections
        ],
    }


def analyser_classeur(fichier) -> dict:
    """Lit le classeur et renvoie client (« Facturer à ») + une entrée par
    feuille de zone, sans rien écrire en base."""
    wb = load_workbook(fichier, data_only=True)
    client = None
    feuilles = []
    for ws in wb.worksheets:
        entete = _ligne_entete(ws)
        if not entete:
            continue
        contenu = _lire_feuille(ws, entete)
        nb = sum(len(s["extincteurs"]) for s in contenu["sections"])
        if not nb:
            continue
        if client is None:
            ville_cp = _texte(ws["D13"].value)
            client = {
                "nom": _texte(ws["D11"].value),
                "adresse": _texte(ws["D12"].value),
                "ville": re.sub(r"\s+,", ",", ville_cp).title().replace(", Qc", ", QC"),
                "code_postal": _texte(ws["D14"].value).upper(),
                "contact_nom": _texte(ws["C15"].value),
                "contact_telephone": _texte(ws["C16"].value),
                "contact_email": _texte(ws["C17"].value),
            }
        ville = _texte(ws["J13"].value).split(",")[0].strip().title()
        feuilles.append({
            "feuille": ws.title,
            "nom": _nom_lieu(ws.title),
            "adresse": _texte(ws["J12"].value),
            "ville": ville,
            "code_postal": _texte(ws["J14"].value).upper(),
            "nb_extincteurs": nb,
            "nb_sections": len(contenu["sections"]),
            "nb_non_conformites": sum(
                1 for s in contenu["sections"] for x in s["extincteurs"] if x["non_conformites"] or x["remarque"]
            ),
            "sections": contenu["sections"],
        })
    return {"client": client or {}, "feuilles": feuilles}


def _adresse_en_parties(adresse: str) -> tuple[str, str]:
    m = re.match(r"^\s*(\d+[A-Za-z]?)\s+(.*)$", adresse or "")
    return (m.group(1), m.group(2).strip()) if m else ("", (adresse or "").strip())
